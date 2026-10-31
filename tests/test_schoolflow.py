import io
import shutil
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from core import schoolformat, docx2tex, texparse, profile as P
from web import app as web


class SchoolRouting(unittest.TestCase):
    def test_reference_roundtrip_has_one_number_and_no_internal_commands(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'source.docx'
            doc = Document()
            for text in ['1 绪论', '正文引用[1]。', '参考文献', '[1] 测试文献. 题目[M]. 2026.']:
                doc.add_paragraph(text)
            doc.save(source)
            tex, _ = docx2tex.convert(str(source), tmp, P.load('hkjjxy'), log=lambda *a: None)
            refs = [b['text'] for b in texparse.parse(tex)['blocks'] if b['kind'] == 'ref']
            self.assertEqual(refs, ['[1] 测试文献. 题目[M]. 2026.'])
            self.assertIn(r'\textsuperscript{[1]}', tex)

    def setUp(self):
        self.client = TestClient(web.app)

    def test_haikou_search_uses_local_template(self):
        data = self.client.get('/api/format/rules', params={'q': '海口经济学院'}).json()
        self.assertEqual(data['items'][0]['id'], 'school:hkjjxy')
        self.assertEqual(data['items'][0]['provider'], 'local')

    def test_other_school_gets_provider_namespace(self):
        with patch.object(web.dayipaper, 'has_token', return_value=True), patch.object(
                web.dayipaper, 'search_rules', return_value={'items': [{'id': 'r'}], 'total': 1}):
            data = self.client.get('/api/format/rules?q=other').json()
            self.assertEqual(data['items'][0]['id'], 'dayipaper:r')

    def test_haikou_submission_uses_local_template(self):
        with patch.object(web, '_run_school') as local:
            response = self.client.post('/api/format', files={'file': ('paper.docx', b'input')},
                data={'ruleId': 'school:hkjjxy', 'title': 'Title', 'authorName': 'Author'})
            self.assertEqual(response.status_code, 200)
            jid = response.json()['job']
            try:
                self.assertEqual(web.JOBS[jid]['kind'], 'school')
                local.assert_called_once()
                self.assertEqual(local.call_args.args[2]['ruleId'], 'school:hkjjxy')
            finally:
                job = web.JOBS.pop(jid)
                shutil.rmtree(job['tmp'])

    def test_unknown_school_cannot_become_a_path(self):
        with patch.object(web.dayipaper, 'has_token', return_value=True):
            response = self.client.post('/api/format', files={'file': ('paper.docx', b'input')},
            data={'ruleId': 'school:../../outside', 'title': 'Title', 'authorName': 'Author'})
        self.assertEqual(response.status_code, 400)

    def test_cover_keeps_template_fields_unfilled(self):
        cfg = P.load('hkjjxy')
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / 'cover.docx'
            schoolformat.prepare_cover(cfg, {'title': '论文测试标题', 'authorName': '测试作者',
                'tutorName': '测试导师', 'subjectName': '测试专业'}, target)
            text = '\n'.join(p.text for p in Document(target).paragraphs)
            self.assertNotIn('论文测试标题', text)
            self.assertNotIn('测试作者', text)
            self.assertNotIn('测试导师', text)
            template = Document(P.path(cfg, cfg['cover']['template']))
            self.assertEqual(text, '\n'.join(p.text for p in template.paragraphs))
