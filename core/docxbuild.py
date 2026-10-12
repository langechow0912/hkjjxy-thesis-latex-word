# -*- coding: utf-8 -*-
"""中间结构 → 可编辑 Word（.docx）。

所有字体 / 字号 / 行距 / 间距都从 profile.yaml 读，这里不出现任何学校专属数值。

三处关键手法（踩过坑才有的）：
  * 三线表：tblBorders 只开 top/bottom，其余置 none；表头行再叠一条 bottom 细线
  * 表题/图题：用 SEQ 域自动编号并写入缓存值，用户改稿后全选 F9 即可重排
  * 目录：TOC 域 + 缓存里预填真实页码（两遍构建，见 toc.py）
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT,
                            WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import profile as P

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT,
          "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT,
          "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

_LEADER = {"dots": WD_TAB_LEADER.DOTS,
           "none": WD_TAB_LEADER.SPACES,
           "hyphen": WD_TAB_LEADER.DASHES}


# ---------------------------------------------------------------- 基础工具
def fmt(run, cn="宋体", en="Times New Roman", size=12, bold=False, sup=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if sup:
        run.font.superscript = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), cn)


def para(doc, align=None, before=0, after=0, ls=1.5, indent_first=None,
         left=None, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = ls
    if indent_first is not None:
        pf.first_line_indent = Pt(indent_first)
    if left is not None:
        pf.left_indent = Pt(left)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    return p


def heading(doc, text, level, spec):
    p = para(doc, before=spec.get("before", 6), after=spec.get("after", 6))
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level - 1))
    pPr.append(ol)
    r = p.add_run(text)
    fmt(r, cn=spec.get("font", "黑体"), size=P.pt(spec.get("size", 12)),
        bold=spec.get("bold", False))
    return p


def add_tc_field(par, text):
    """给段落挂一个 TC 域，让「摘要/关键词」这类非标题段落也能进目录"""
    r1 = par.add_run()._r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TC "%s" \\f C \\l 1 ' % text
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    for el in (b, it, e):
        r1.append(el)


# ---------------------------------------------------------------- 图表题
def _seq_run(doc, n, seq_name, size, east_asia="黑体", bold=True):
    """插入 ` SEQ 表|图 ` 域，并把当前编号写进缓存文本（不更新域也能看到）"""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r" SEQ %s \* ARABIC " % seq_name)
    rr = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:eastAsia"), east_asia)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(rf)
    rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:b"))
    t = OxmlElement("w:t")
    t.text = str(n)
    rr.append(rpr)
    rr.append(t)
    fld.append(rr)
    return fld


def table_caption(doc, n, name, cap_cfg):
    p = para(doc, align=_ALIGN[cap_cfg.get("align", "center")],
             before=cap_cfg.get("before", 4), after=cap_cfg.get("after", 2),
             ls=cap_cfg.get("linespread", 1.0))
    size = P.pt(cap_cfg.get("size", 10.5))
    cn = cap_cfg.get("font", "黑体")
    r1 = p.add_run("表")
    fmt(r1, cn=cn, size=size, bold=True)
    p._p.append(_seq_run(doc, n, "表", size, east_asia=cn))
    r2 = p.add_run("\u3000" + name)
    fmt(r2, cn=cn, size=size, bold=True)
    return p


def figure_caption(doc, n, name, cap_cfg):
    p = para(doc, align=_ALIGN[cap_cfg.get("align", "center")],
             before=cap_cfg.get("before", 2), after=cap_cfg.get("after", 6),
             ls=cap_cfg.get("linespread", 1.0))
    size = P.pt(cap_cfg.get("size", 10.5))
    cn = cap_cfg.get("font", "黑体")
    r1 = p.add_run("图")
    fmt(r1, cn=cn, size=size, bold=True)
    p._p.append(_seq_run(doc, n, "图", size, east_asia=cn))
    r2 = p.add_run("\u3000" + name)
    fmt(r2, cn=cn, size=size, bold=True)
    return p


def insert_figure(doc, fig_dir, rel_path, frac, cfg):
    img = os.path.join(fig_dir, os.path.basename(rel_path))
    if not os.path.isfile(img):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run("［缺图：%s］" % rel_path)
        fmt(r, size=10.5)
        return p
    doc.add_picture(img, width=Pt(P.textwidth_pt(cfg) * frac))
    # add_picture 返回 InlineShape；承载它的段落是最后一段
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return p


def three_line_table(doc, rows, cfg):
    tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tcfg = cfg["table"]
    top_sz = int(round(float(tcfg.get("toprule_bp", 1.5)) * 8))   # w:sz 单位 = 1/8 pt
    mid_sz = int(round(float(tcfg.get("midrule_bp", 0.5)) * 8))

    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "bottom"):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(top_sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for tag in ("left", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tb._tbl.tblPr.append(borders)

    bsize = P.pt(tcfg["body"].get("size", 10.5))
    bls = tcfg["body"].get("linespread", 1.0)
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = tb.cell(ri, ci)
            p = cell.paragraphs[0]
            p.alignment = _ALIGN[tcfg["body"].get("first_col_align", "left")] \
                if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            pf.line_spacing = bls
            r = p.add_run(cell_text)
            fmt(r, size=bsize)
            if ri == 0:  # 栏目线
                tcPr = cell._tc.get_or_add_tcPr()
                tcb = OxmlElement("w:tcBorders")
                bt = OxmlElement("w:bottom")
                bt.set(qn("w:val"), "single")
                bt.set(qn("w:sz"), str(mid_sz))
                bt.set(qn("w:color"), "000000")
                tcb.append(bt)
                tcPr.append(tcb)
    return tb


def body_paragraph(doc, text, cfg):
    """正文段落；把 \\textsuperscript{[n]} 渲染成上标"""
    if not text:
        return None
    p = para(doc, indent_first=P.pt(cfg["body"]["size"]) * cfg["body"]["first_indent"])
    size = P.pt(cfg["body"]["size"])
    for part in re.split(r"(\\textsuperscript\{[^}]*\})", text):
        m = re.match(r"\\textsuperscript\{([^}]*)\}", part)
        if m:
            r = p.add_run(m.group(1))
            fmt(r, size=size, sup=True)
        elif part:
            r = p.add_run(part)
            fmt(r, size=size)
    return p


def page_break(doc):
    pbrk = doc.add_paragraph()
    pf = pbrk.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pbrk.add_run().add_break(WD_BREAK.PAGE)
    return pbrk


# ---------------------------------------------------------------- 目录样式
def define_toc_styles(doc, cfg):
    tcfg = cfg["toc"]
    for lv in (1, 2, 3):
        nm = "toc %d" % lv
        try:
            st = doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            st = doc.styles[nm]
        st.base_style = doc.styles["Normal"]
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Pt(0)
        pf.first_line_indent = Pt(0)
        pf.line_spacing = tcfg.get("linespread", 1.5)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        f = st.font
        f.name = "Times New Roman"
        f.size = Pt(P.pt(tcfg.get("entry_size", 12)))
        st.element.rPr.rFonts.set(qn("w:eastAsia"), tcfg.get("entry_font", "宋体"))


def toc_entry_par(doc, cfg, first=False, last=False, entry=None,
                  placeholder=None, instr=""):
    """一个目录条目；域字符跨段落包裹整组条目（begin…separate … end）"""
    p = para(doc, ls=cfg["toc"].get("linespread", 1.5))
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(
        Pt(P.textwidth_pt(cfg)), WD_TAB_ALIGNMENT.RIGHT,
        _LEADER[cfg["toc"].get("leader", "dots")])
    if first:
        r = p.add_run()._r
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        for el in (b, it, sep):
            r.append(el)
    if entry is None:
        r = p.add_run(placeholder)
        fmt(r, size=P.pt(cfg["toc"].get("entry_size", 12)))
    else:
        text, page = entry
        size = P.pt(cfg["toc"].get("entry_size", 12))
        r = p.add_run(text)
        fmt(r, size=size)
        r2 = p.add_run()
        r2.add_tab()
        fmt(r2, size=size)
        r3 = p.add_run(str(page))
        fmt(r3, size=size)
    if last:
        r = p.add_run()._r
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        r.append(e)
    return p


def add_page_number_footer(section, cfg):
    """页码：按 profile 的 page.pagenum 决定放页脚还是页眉、居中还是靠右"""
    pcfg = cfg["page"]
    pn = pcfg.get("pagenum", {"position": "footer", "align": "center"})
    container = section.footer if pn.get("position", "footer") == "footer" \
        else section.header
    container.is_linked_to_previous = False
    p = container.paragraphs[0]
    p.alignment = _ALIGN[pn.get("align", "center")]
    r = p.add_run()._r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = r" PAGE \* ARABIC "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    for el in (b, it, sep):
        r.append(el)
    rt = p.add_run("1")
    fmt(rt, size=10.5)
    p.add_run()._r.append(e)


# ---------------------------------------------------------------- 主构建
def build(cfg, struct, out_path, fig_dir, tpl_path, entries_map=None):
    """把解析结果写成 docx。

    entries_map 为 None 时是「探针」pass：目录只放占位文本，用来渲染取页码。
    """
    doc = Document(tpl_path)
    define_toc_styles(doc, cfg)

    page = cfg["page"]
    m = page["margins"]

    def setup_section(s):
        s.page_width = Pt(P.mm(page["width"]))
        s.page_height = Pt(P.mm(page["height"]))
        s.top_margin = Pt(P.mm(m["top"]))
        s.bottom_margin = Pt(P.mm(m["bottom"]))
        s.left_margin = Pt(P.mm(m["left"]))
        s.right_margin = Pt(P.mm(m["right"]))

    # ---------------- 目录节 ----------------
    sec_toc = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(sec_toc)
    sec_toc.footer.is_linked_to_previous = False

    tcfg = cfg["toc"]
    depth = int(tcfg.get("depth", 3))
    specs = [("toc", tcfg.get("title", "目　录"),
              r' TOC \o "1-%d" \h \z \u \f C ' % depth)]
    if struct["tables"]:
        specs.append(("lot", cfg["table"].get("list_title", "表目录"),
                      r' TOC \h \z \c "表" '))
    if struct["figures"]:
        specs.append(("lof", cfg["figure"].get("list_title", "图目录"),
                      r' TOC \h \z \c "图" '))

    for idx, (key, title, instr) in enumerate(specs):
        if idx == 0:
            p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12)
            r = p.add_run(title)
            fmt(r, size=P.pt(tcfg.get("title_size", 18)), bold=True)
        else:
            p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=12)
            r = p.add_run(title)
            fmt(r, size=P.pt(tcfg.get("title_size", 18)), bold=True)
        data = (entries_map or {}).get(key)
        items = data if data else [None]
        for i, e in enumerate(items):
            toc_entry_par(doc, cfg, first=(i == 0), last=(i == len(items) - 1),
                          entry=e, placeholder="（自动目录域：全选后按 F9 更新）",
                          instr=instr)

    # ---------------- 正文节（页码重新从 1 起） ----------------
    sec_body = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(sec_body)
    sec_body.footer.is_linked_to_previous = False
    pgnum = OxmlElement("w:pgNumType")
    pgnum.set(qn("w:start"), str(cfg["page"].get("body_page_start", 1)))
    sectPr = sec_body._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pgnum)
    else:
        sectPr.append(pgnum)
    add_page_number_footer(sec_body, cfg)

    meta = struct["meta"]
    _write_front_matter(doc, cfg, meta)
    _write_blocks(doc, cfg, struct["blocks"], fig_dir)

    doc.save(out_path)
    return out_path


def _write_front_matter(doc, cfg, meta):
    """中文题目 / 中英文摘要关键词"""
    ab = cfg["abstract"]
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0)
    r = p.add_run(meta["cn_title"])
    fmt(r, cn=ab.get("title_font", "黑体"), size=P.pt(ab.get("title_size", 18)),
        bold=True)
    if meta.get("cn_subtitle"):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
        r = p.add_run(meta["cn_subtitle"])
        fmt(r, cn=ab.get("title_font", "黑体"), size=P.pt(ab.get("title_size", 18)),
            bold=True)

    p = para(doc, after=4)
    r = p.add_run("摘\u3000\u3000要：")
    fmt(r, cn=ab.get("font", "楷体"), size=P.pt(ab.get("label_size", 12)), bold=True)
    r = p.add_run(meta["cn_abstract"])
    fmt(r, cn=ab.get("font", "楷体"), size=P.pt(ab.get("body_size", 10.5)))
    add_tc_field(p, "摘要")

    p = para(doc, after=6)
    r = p.add_run("关键词：")
    fmt(r, cn=ab.get("font", "楷体"), size=P.pt(ab.get("label_size", 12)), bold=True)
    r = p.add_run(meta["cn_keywords"])
    fmt(r, cn=ab.get("font", "楷体"), size=P.pt(ab.get("body_size", 10.5)))
    add_tc_field(p, "关键词")

    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6)
    r = p.add_run(meta["en_title"])
    fmt(r, size=P.pt(ab.get("en_title_size", 15)), bold=True)

    p = para(doc, after=4)
    r = p.add_run("Abstract: ")
    fmt(r, size=P.pt(ab.get("body_size", 10.5)), bold=True)
    r = p.add_run(meta["en_abstract"])
    fmt(r, size=P.pt(ab.get("body_size", 10.5)))

    p = para(doc, after=6)
    r = p.add_run("Key words: ")
    fmt(r, size=P.pt(ab.get("body_size", 10.5)), bold=True)
    r = p.add_run(meta["en_keywords"])
    fmt(r, size=P.pt(ab.get("body_size", 10.5)))


def _write_blocks(doc, cfg, blocks, fig_dir):
    tab_no = [0]
    fig_no = [0]
    for b in blocks:
        kind = b["kind"]
        if kind == "break":
            page_break(doc)
        elif kind == "heading":
            heading(doc, b["text"], b["level"], cfg["heading"]["levels"][b["level"]])
        elif kind == "unnum":
            heading(doc, b["text"], 1, cfg["heading"]["unnum"])
        elif kind == "subheading":
            p = para(doc, before=6, after=3,
                     indent_first=P.pt(cfg["body"]["size"]) * cfg["body"]["first_indent"])
            r = p.add_run(b["text"])
            fmt(r, size=P.pt(cfg["body"]["size"]), bold=True)
        elif kind == "table":
            tab_no[0] += 1
            if cfg["table"]["caption"].get("position", "above") == "above":
                table_caption(doc, tab_no[0], b["name"], cfg["table"]["caption"])
                three_line_table(doc, b["rows"], cfg)
            else:
                three_line_table(doc, b["rows"], cfg)
                table_caption(doc, tab_no[0], b["name"], cfg["table"]["caption"])
            if b.get("note"):
                p = para(doc, before=2, after=6, ls=1.0)
                r = p.add_run(b["note"])
                fmt(r, size=P.pt(cfg["table"]["note"]["size"]))
        elif kind == "figure":
            fig_no[0] += 1
            insert_figure(doc, fig_dir, b["path"], b.get("frac", 0.82), cfg)
            figure_caption(doc, fig_no[0], b["caption"], cfg["figure"]["caption"])
        elif kind == "ref":
            p = para(doc, ls=1.2, left=24, hanging=24)
            r = p.add_run(b["text"])
            fmt(r, size=P.pt(cfg["reference"]["size"]))
        elif kind == "text":
            body_paragraph(doc, b["text"], cfg)


def toc_entry_list(cfg, struct):
    """目录条目列表 [(level, 显示文本)…]。

    摘要/关键词这类非 \\section 的块由 profile 的 toc.front_entries 声明。
    """
    depth = int(cfg["toc"].get("depth", 3))
    front = [(1, t) for t in cfg["toc"].get("front_entries", [])]
    rest = [(lv, t) for lv, t in struct["toc"] if lv <= depth]
    return front + rest
