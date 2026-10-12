# -*- coding: utf-8 -*-
"""从 main.tex 结构生成可编辑 Word 版（海口经济学院毕业论文格式）
两遍构建：pass1 生成探针 → LibreOffice 渲染取页码 → pass2 预填目录
"""
import io, re, os, subprocess, glob
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEX = io.open(r"C:/Users/周良/Downloads/排版测试论文/以岭药业-hkjjxy/main.tex", encoding="utf-8").read()
TPL = r"C:\Users\周良\Desktop\海口经济学院模板前2页.docx"
FIG_DIR = r"C:/Users/周良/Downloads/排版测试论文/以岭药业-hkjjxy/figures"
OUT_DIR = r"C:/Users/周良/Downloads/排版测试论文/以岭药业-hkjjxy"
OUT = OUT_DIR + "/以岭药业财务风险分析-毕业论文（可编辑）v2.docx"
TMP = os.path.dirname(os.path.abspath(__file__))
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

def grab1(pat):
    m = re.search(pat, TEX)
    return m.group(1).strip() if m else ""

CN_TITLE = grab1(r"\\cnTitle\{([^}]*)\}")
CN_SUB   = grab1(r"\\cnSubTitle\{([^}]*)\}")
EN_TITLE = grab1(r"\\enTitle\{([^}]*)\}")
CN_ABS   = grab1(r"\\cnAbstract\{%([\s\S]*?)\n\}")
EN_ABS   = grab1(r"\\enAbstract\{%([\s\S]*?)\n\}")
CN_KEY   = grab1(r"\\cnKeywords\{([^}]*)\}")
EN_KEY   = grab1(r"\\enKeywords\{([^}]*)\}")

def unesc(s):
    s = s.replace(r"\&", "&").replace(r"\%", "%").replace(r"\#", "#")
    s = s.replace(r"\_", "_").replace(r"\{", "{").replace(r"\}", "}").replace("~", " ")
    return re.sub(r"\s+", " ", s).strip()

CN_ABS, EN_ABS = unesc(CN_ABS), unesc(EN_ABS)
CN_KEY, EN_KEY = unesc(CN_KEY), unesc(EN_KEY)

# ---------- 字体/段落工具 ----------
def fmt(run, cn="宋体", en="Times New Roman", size=12, bold=False, sup=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if sup:
        run.font.superscript = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), cn)

def para(doc, align=None, before=0, after=0, ls=1.5, indent_first=None,
         left=None, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = ls
    if indent_first is not None:
        pf.first_line_indent = Pt(indent_first)
    if left is not None:
        pf.left_indent = Pt(left)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    return p

def heading(doc, text, level, cn="黑体", size=15, before=12, after=6):
    p = para(doc, before=before, after=after)
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(level - 1))
    pPr.append(ol)
    r = p.add_run(text)
    fmt(r, cn=cn, size=size)
    return p

def add_tc_field(par, text):
    r1 = par.add_run()._r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = f' TC "{text}" \\f C \\l 1 '
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, it, e):
        r1.append(el)

def seq_caption(doc, n, name):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, ls=1.0)
    r1 = p.add_run("表")
    fmt(r1, cn="黑体", size=10.5, bold=True)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r" SEQ 表 \* ARABIC ")
    rr = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:eastAsia"), "黑体")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21")
    bold = OxmlElement("w:b")
    rpr.append(rf); rpr.append(sz); rpr.append(bold)
    t = OxmlElement("w:t"); t.text = str(n)
    rr.append(rpr); rr.append(t)
    fld.append(rr)
    p._p.append(fld)
    r2 = p.add_run("\u3000" + name)
    fmt(r2, cn="黑体", size=10.5, bold=True)

def fig_caption(doc, n, name):
    """图题：图 N　题名，黑体五号加粗，居中，位于图下方"""
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=6, ls=1.0)
    r1 = p.add_run("图")
    fmt(r1, cn="黑体", size=10.5, bold=True)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r" SEQ 图 \* ARABIC ")
    rr = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts"); rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:eastAsia"), "黑体")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "21")
    bold = OxmlElement("w:b")
    rpr.append(rf); rpr.append(sz); rpr.append(bold)
    t = OxmlElement("w:t"); t.text = str(n)
    rr.append(rpr); rr.append(t)
    fld.append(rr)
    p._p.append(fld)
    r2 = p.add_run("\u3000" + name)
    fmt(r2, cn="黑体", size=10.5, bold=True)

def insert_figure(doc, rel_path, frac):
    """插入图片（版心宽 165mm × frac），居中"""
    img = os.path.join(FIG_DIR, os.path.basename(rel_path))
    width = Mm(165 * frac)
    doc.add_picture(img, width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = doc.paragraphs[-1].paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(0); pf.line_spacing = 1.0

def three_line_table(doc, rows):
    tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tb._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, sz in (("top", 12), ("bottom", 12), ("left", None), ("right", None),
                    ("insideH", None), ("insideV", None)):
        el = OxmlElement(f"w:{tag}")
        if sz:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "none"); el.set(qn("w:sz"), "0"); el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = tb.cell(ri, ci)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.0
            r = p.add_run(cell_text)
            fmt(r, size=10.5)
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tcb = OxmlElement("w:tcBorders")
                bt = OxmlElement("w:bottom")
                bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4"); bt.set(qn("w:color"), "000000")
                tcb.append(bt); tcPr.append(tcb)

# ---------- 预扫：表格与标签 ----------
lines = TEX.split("\n")
i0 = lines.index("\\begin{document}")

tab_labels, tab_total = {}, 0
tables = []
k = 0
while k < len(lines):
    ln = lines[k]
    if "\\begin{thesTab}" in ln:
        m = re.search(r"\{([^{}]*?)\\label\{(tab:t\d+)\}\}\{([^{}]*)\}\s*$", ln)
        tname, tlab, tnote = m.group(1).strip(), m.group(2), m.group(3).strip()
        tab_total += 1
        tab_labels[tlab.split(":")[1]] = tab_total
        rows = []
        k += 1
        while "\\end{thesTab}" not in lines[k]:
            l2 = lines[k].strip()
            if l2.endswith("\\\\") and not l2.startswith("\\"):
                rows.append([unesc(c.strip()) for c in l2[:-2].split("&")])
            k += 1
        tables.append((tname, tnote, rows))
    k += 1

# ---------- 预扫：图环境 ----------
figures = []   # (rel_path, frac, caption)
fig_total = 0
k = 0
while k < len(lines):
    if "\\begin{figure}" in lines[k]:
        blk = []
        k += 1
        while "\\end{figure}" not in lines[k]:
            blk.append(lines[k]); k += 1
        blk = "\n".join(blk)
        mi = re.search(r"includegraphics\[([^\]]*)\]\{([^}]*)\}", blk)
        mc = re.search(r"caption\{([^{}]*)\}", blk)
        frac = 0.82
        if mi:
            mw = re.search(r"width=([\d.]+)\\linewidth", mi.group(1))
            if mw:
                frac = float(mw.group(1))
        figures.append((mi.group(2) if mi else "", frac, mc.group(1).strip() if mc else ""))
        fig_total += 1
    k += 1

def resolve(s):
    return re.sub(r"~?\\ref\{tab:(t\d+)\}",
                  lambda m: "表" + str(tab_labels.get(m.group(1), "?")), s)

def add_body_text(doc, text):
    text = resolve(unesc(text))
    if not text:
        return
    p = para(doc, indent_first=24)
    for part in re.split(r"(\\textsuperscript\{[^}]*\})", text):
        m = re.match(r"\\textsuperscript\{([^}]*)\}", part)
        if m:
            r = p.add_run(m.group(1)); fmt(r, size=12, sup=True)
        elif part:
            r = p.add_run(part); fmt(r, size=12)

# ---------- 目录条目预扫描（收集 level + 显示文本，页码第二遍回填） ----------
toc_entries = []   # (level, display_text)
lot_entries = []   # 表名

# ---------- 构建 ----------
def define_toc_styles(doc):
    for nm in ("toc 1", "toc 2", "toc 3"):
        try:
            st = doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            st = doc.styles[nm]
        st.base_style = doc.styles["Normal"]
        pf2 = st.paragraph_format
        pf2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf2.left_indent = Pt(0); pf2.first_line_indent = Pt(0)
        pf2.line_spacing = 1.5
        pf2.space_before = Pt(0); pf2.space_after = Pt(0)
        f2 = st.font
        f2.name = "Times New Roman"; f2.size = Pt(12)
        f2.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

TAB_POS = Pt((210 - 25 - 20) / 25.4 * 72)   # 版心宽 165mm

def toc_entry_par(doc, first=False, last=False, entry=None, placeholder=None):
    """一个目录条目段落；域字符跨段落包裹所有条目"""
    p = para(doc, ls=1.5)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(TAB_POS, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if first:
        r = p.add_run()._r
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = CURRENT_INSTR
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        for el in (b, it, sep):
            r.append(el)
    if entry is None:
        rt = p.add_run(placeholder); fmt(rt, size=12)
    else:
        text, page = entry
        rt = p.add_run(text); fmt(rt, size=12)
        rt2 = p.add_run(); rt2.add_tab()
        fmt(rt2, size=12)
        rp = p.add_run(str(page)); fmt(rp, size=12)
    if last:
        re_ = p.add_run()._r
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        re_.append(e)
    return p

CURRENT_INSTR = ""

def build(entries_map, path):
    """entries_map: {"toc": [(text,page)...] or None, "lot": [...] or None}"""
    global CURRENT_INSTR
    doc = Document(TPL)
    define_toc_styles(doc)

    def setup_section(s):
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = Mm(30)
        s.left_margin, s.right_margin = Mm(25), Mm(20)

    sec_toc = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(sec_toc)
    sec_toc.footer.is_linked_to_previous = False

    # ================= 目录页 =================
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12)
    r = p.add_run("目\u3000\u3000\u3000录"); fmt(r, size=18, bold=True)

    CURRENT_INSTR = r' TOC \o "1-3" \h \z \u \f C '
    data = entries_map.get("toc") if entries_map else None
    if data:
        for i, e in enumerate(data):
            toc_entry_par(doc, first=(i == 0), last=(i == len(data) - 1), entry=e)
    else:
        toc_entry_par(doc, first=True, last=True,
                      placeholder="（目录为自动域：全选后按 F9 或右键“更新域”生成）")

    # ---- 表目录（正文中有表才加）----
    if tab_total > 0:
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=12)
        r = p.add_run("表\u3000\u3000目\u3000\u3000录"); fmt(r, size=18, bold=True)
        CURRENT_INSTR = r' TOC \h \z \c "表" '
        data = entries_map.get("lot") if entries_map else None
        if data:
            for i, e in enumerate(data):
                toc_entry_par(doc, first=(i == 0), last=(i == len(data) - 1), entry=e)
        else:
            toc_entry_par(doc, first=True, last=True,
                          placeholder="（表目录为自动域：按 F9 更新生成）")

    # ---- 图目录（正文中有图才加）----
    has_fig = "\\begin{figure}" in TEX
    if has_fig:
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=12)
        r = p.add_run("图\u3000\u3000目\u3000\u3000录"); fmt(r, size=18, bold=True)
        CURRENT_INSTR = r' TOC \h \z \c "图" '
        data = entries_map.get("lof") if entries_map else None
        if data:
            for i, e in enumerate(data):
                toc_entry_par(doc, first=(i == 0), last=(i == len(data) - 1), entry=e)
        else:
            toc_entry_par(doc, first=True, last=True,
                          placeholder="（图目录为自动域：按 F9 更新生成）")

    # ================= 正文节（页码从 1 起） =================
    sec_body = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_section(sec_body)
    sec_body.footer.is_linked_to_previous = False
    pgnum = OxmlElement("w:pgNumType"); pgnum.set(qn("w:start"), "1")
    sectPr = sec_body._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pgnum)
    else:
        sectPr.append(pgnum)
    fp = sec_body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 页码域
    rb = fp.add_run()._r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = r" PAGE \* ARABIC "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    for el in (b, it, sep):
        rb.append(el)
    rt = fp.add_run("1"); fmt(rt, size=10.5)
    re_ = fp.add_run()._r
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    re_.append(e)

    # ---- 中文题目 ----
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0)
    r = p.add_run(CN_TITLE); fmt(r, cn="黑体", size=18)
    if CN_SUB:
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
        r = p.add_run(CN_SUB); fmt(r, cn="黑体", size=18)

    # ---- 中文摘要/关键词 ----
    p = para(doc, after=4)
    r = p.add_run("摘\u3000\u3000要："); fmt(r, cn="楷体", size=12, bold=True)
    r = p.add_run(CN_ABS); fmt(r, cn="楷体", size=10.5)
    add_tc_field(p, "摘要")
    toc_entries.append((1, "摘要"))

    p = para(doc, after=6)
    r = p.add_run("关键词："); fmt(r, cn="楷体", size=12, bold=True)
    r = p.add_run(CN_KEY); fmt(r, cn="楷体", size=10.5)
    add_tc_field(p, "关键词")
    toc_entries.append((1, "关键词"))

    # ---- 英文题目/摘要 ----
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6)
    r = p.add_run(EN_TITLE); fmt(r, size=15, bold=True)

    p = para(doc, after=4)
    r = p.add_run("Abstract: "); fmt(r, size=10.5, bold=True)
    r = p.add_run(EN_ABS); fmt(r, size=10.5)

    p = para(doc, after=6)
    r = p.add_run("Key words: "); fmt(r, size=10.5, bold=True)
    r = p.add_run(EN_KEY); fmt(r, size=10.5)

    # ---- 逐行扫描正文 ----
    cnt = [0, 0, 0]
    buf = []
    refn = 0
    tab_iter = iter(tables)
    fig_iter = iter(figures)
    fig_no = [0]

    def flush():
        if buf:
            add_body_text(doc, " ".join(buf))
            buf.clear()

    SKIP = ("\\setcounter", "\\pagenumbering", "\\pagestyle", "\\cnTitleBlock",
            "\\enTitleBlock", "\\includepdf", "\\phantomsection", "\\addcontentsline",
            "\\documentclass", "\\usepackage", "\\cnTitle{", "\\cnSubTitle{", "\\enTitle{",
            "\\begin{document}", "\\end{document}", "\\toprule", "\\midrule", "\\bottomrule",
            "\\begin{tabularx}", "\\end{tabularx}", "\\begin{table}", "\\end{table}",
            "\\vspace", "\\tocpage", "\\cnKeywords{", "\\enKeywords{", "\\listoffigures",
            "\\listoftables", "\\tableofcontents")

    k = i0 + 1
    while k < len(lines):
        ln = lines[k].strip()
        if ln.startswith("%") or ln == "":
            flush(); k += 1; continue

        if ln.startswith("\\newpage"):
            # 【对应 LaTeX \newpage】正文分页：摘要→引言、结论后→参考文献 等处另起一页
            flush()
            pbrk = doc.add_paragraph()
            pbrk.paragraph_format.space_before = Pt(0)
            pbrk.paragraph_format.space_after = Pt(0)
            pbrk.paragraph_format.line_spacing = 1.0
            r = pbrk.add_run()
            r.add_break(WD_BREAK.PAGE)
            rest = ln[len("\\newpage"):].strip()
            if rest:
                lines[k] = rest
                continue
            k += 1; continue

        if "\\begin{thesTab}" in ln:
            flush()
            tname, tnote, rows = next(tab_iter)
            seq_caption(doc, tab_total - len(tables), tname)
            three_line_table(doc, rows)
            if tnote:
                p = para(doc, before=2, after=6, ls=1.0)
                r = p.add_run(unesc(tnote)); fmt(r, size=9)
            while "\\end{thesTab}" not in lines[k]:
                k += 1
            k += 1; continue

        if "\\begin{figure}" in ln:
            flush()
            fig_no[0] += 1
            rel, frac, cap = next(fig_iter)
            insert_figure(doc, rel, frac)
            fig_caption(doc, fig_no[0], cap)
            while "\\end{figure}" not in lines[k]:
                k += 1
            k += 1; continue

        m = re.match(r"\\section(\*)?\{(.+)\}\s*$", ln)
        if m:
            flush()
            title = resolve(unesc(m.group(2)))
            if m.group(1):
                heading(doc, title, 1)
                toc_entries.append((1, title))
            else:
                cnt[0] += 1; cnt[1] = cnt[2] = 0
                t = f"{cnt[0]}\u3000{title}"
                heading(doc, t, 1)
                toc_entries.append((1, t))
            k += 1; continue
        m = re.match(r"\\subsection\{(.+)\}\s*$", ln)
        if m:
            flush()
            cnt[1] += 1; cnt[2] = 0
            t = f"{cnt[0]}.{cnt[1]}\u3000{resolve(unesc(m.group(1)))}"
            heading(doc, t, 2, size=14, before=8, after=4)
            toc_entries.append((2, t))
            k += 1; continue
        m = re.match(r"\\subsubsection\{(.+)\}\s*$", ln)
        if m:
            flush()
            cnt[2] += 1
            t = f"{cnt[0]}.{cnt[1]}.{cnt[2]}\u3000{resolve(unesc(m.group(1)))}"
            heading(doc, t, 3, size=12, before=6, after=3)
            toc_entries.append((3, t))
            k += 1; continue

        m = re.match(r"\\paragraph\*?\{(.+)\}\s*$", ln)
        if m:
            flush()
            p = para(doc, before=6, after=3, indent_first=24)
            r = p.add_run(resolve(unesc(m.group(1)))); fmt(r, size=12, bold=True)
            k += 1; continue

        if ln.startswith("\\unnumsec{"):
            flush()
            t = resolve(unesc(ln[len("\\unnumsec{"):-1]))
            heading(doc, t, 1, size=12, before=54, after=6)
            toc_entries.append((1, t))
            k += 1; continue
        if "\\begin{thesref}" in ln or "\\end{thesref}" in ln:
            flush(); k += 1; continue
        if ln.startswith("\\item "):
            refn += 1
            p = para(doc, ls=1.2, left=24, hanging=24)
            r = p.add_run(f"[{refn}] " + resolve(unesc(ln[6:]))); fmt(r, size=10.5)
            k += 1; continue
        # 中英文摘要块为多行结构（\cnAbstract{% ... 换行 ... }），整块跳过，
        # 摘要正文由脚本前段按格式规范显式写入，不能在此重复输出
        if ln.startswith(("\\cnAbstract{", "\\enAbstract{")):
            flush()
            while k < len(lines) and lines[k].rstrip() != "}":
                k += 1
            k += 1; continue
        if any(ln.startswith(s) for s in SKIP):
            flush(); k += 1; continue

        buf.append(ln)
        k += 1
    flush()

    doc.save(path)
    return path

# ---------- 页码探测 ----------
def probe_pages():
    probe = os.path.join(TMP, "probe.docx")
    build(None, probe)
    for f in glob.glob(os.path.join(TMP, "probe.pdf")):
        os.remove(f)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", TMP, probe],
                   capture_output=True, timeout=300)
    import fitz
    d = fitz.open(os.path.join(TMP, "probe.pdf"))
    texts = [d[i].get_text().replace("\u3000", " ") for i in range(d.page_count)]
    body_start = next(i for i, t in enumerate(texts) if "Key words" in t)

    cur = body_start
    def find_page(pattern):
        nonlocal cur
        for i in range(cur, len(texts)):
            if re.search(pattern, texts[i]):
                cur = i
                return i - body_start + 1
        return 1

    toc_map, lot_map, lof_map = [], [], []
    for lvl, disp in toc_entries:
        plain = disp.replace("\u3000", " ")
        m = re.match(r"([\d.]+)\s+(.*)", plain)
        if m:
            pat = re.escape(m.group(1)) + r"\s*" + re.escape(m.group(2)[:8])
        else:
            pat = re.escape(plain[:6])
        if plain in ("参考文献",):
            pat = r"参考文献"
        if plain in ("致谢",):
            pat = r"致\s*谢"
        pg = find_page(pat)
        toc_map.append((disp, pg))
    cur = body_start
    for n, tname in enumerate(lot_src, 1):
        pg = find_page(re.escape(tname[:8]))
        lot_map.append((f"表{n}\u3000{tname}", pg))
    cur = body_start
    for n, (rel, frac, cap) in enumerate(figures, 1):
        pg = find_page(r"图\s*" + str(n) + r"\s*[\u3000\s]" + re.escape(cap[:6]))
        lof_map.append((f"图{n}\u3000{cap}", pg))
    d.close()
    return {"toc": toc_map, "lot": lot_map if tab_total > 0 else [],
            "lof": lof_map if fig_total > 0 else []}

# 先收集 lot 源（表名顺序）
lot_src = []
for k_ in range(len(lines)):
    if "\\begin{thesTab}" in lines[k_]:
        m = re.search(r"\{([^{}]*?)\\label\{(tab:t\d+)\}\}", lines[k_])
        lot_src.append(m.group(1).strip())

# ---- pass 1：探针取页码 ----
toc_entries.clear()
probe_map = probe_pages()
print("页码探测完成：", len(probe_map["toc"]), "条目录 /", len(probe_map["lot"]), "条表目录")

# ---- pass 2：正式构建（预填目录） ----
toc_entries.clear()
out = build(probe_map, OUT)
print("saved:", out)

# 清理探针
for f in glob.glob(os.path.join(TMP, "probe.*")):
    os.remove(f)
